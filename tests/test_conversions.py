"""Real-engine integration tests: generate fixtures; validate readable outputs and originals."""

import importlib.util, json, os, subprocess, sys, hashlib, shutil
from pathlib import Path
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "worker"))
import convert as worker

ENGINES = {
    k: str(ROOT / v["path"])
    for k, v in json.loads(
        (ROOT / "engines.json").read_text(encoding="utf-8-sig")
    ).items()
}


@pytest.fixture(scope="session")
def samples():
    from PIL import Image, ImageDraw
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    import ezdxf

    folder = ROOT / "tests" / "fixtures"
    folder.mkdir(exist_ok=True)
    im = Image.new("RGB", (320, 240), "#192331")
    ImageDraw.Draw(im).text((30, 80), "CONVERSOR - teste 2026", fill="#70e0bc")
    for ext in ("png", "jpg", "webp", "bmp", "tiff", "gif"):
        im.save(folder / f"imagem.{ext}")
    im.save(
        folder / "duas-paginas.pdf",
        save_all=True,
        append_images=[Image.new("RGB", (320, 240), "white")],
    )
    doc = Document()
    doc.add_heading("Conversor", 0)
    doc.add_paragraph("Teste de documento: ação e conversão.")
    doc.save(folder / "documento.docx")
    book = Workbook()
    book.active.append(["Produto", "Valor"])
    book.active.append(["Teste", 123.45])
    book.save(folder / "planilha.xlsx")
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[0])
    slide.shapes.title.text = "Conversor"
    slide.placeholders[1].text = "Teste de apresentação"
    ppt.save(folder / "slides.pptx")
    (folder / "texto.txt").write_text(
        "Conversor\nTeste de texto com acentos: ação.", encoding="utf-8"
    )
    (folder / "tabela.csv").write_text("nome,valor\nteste,123\n", encoding="utf-8")
    drawing = ezdxf.new()
    model = drawing.modelspace()
    model.add_line((0, 0), (100, 100))
    model.add_circle((50, 50), 20)
    drawing.saveas(folder / "desenho.dxf")
    shutil.copyfile(
        Path(os.environ["WINDIR"]) / "Fonts" / "arial.ttf", folder / "fonte.ttf"
    )
    worker.run(
        [
            ENGINES["ffmpeg"],
            "-y",
            "-f",
            "lavfi",
            "-i",
            "color=c=blue:s=320x240:r=10",
            "-f",
            "lavfi",
            "-i",
            "sine=frequency=440:sample_rate=44100",
            "-t",
            "0.5",
            "-c:v",
            "libx264",
            "-pix_fmt",
            "yuv420p",
            "-c:a",
            "aac",
            folder / "video.mp4",
        ]
    )
    worker.run(
        [
            ENGINES["ffmpeg"],
            "-y",
            "-i",
            folder / "video.mp4",
            "-vn",
            folder / "audio.wav",
        ]
    )
    worker.run(
        [
            ENGINES["pandoc"],
            folder / "texto.txt",
            "--from",
            "markdown",
            "-o",
            folder / "livro.epub",
        ]
    )
    return folder


def execute(path, target, out, **opts):
    request = dict(
        inputs=[str(path)] if isinstance(path, Path) else [str(x) for x in path],
        target=target,
        outputDirectory=str(out),
        quality=85,
        pages="",
        width=0,
        engines=ENGINES,
        **opts,
    )
    before = {
        p: hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in request["inputs"]
    }
    p = subprocess.run(
        [sys.executable, str(ROOT / "worker" / "convert.py")],
        input=json.dumps(request),
        text=True,
        encoding="utf-8",
        capture_output=True,
        timeout=180,
    )
    assert p.returncode == 0, p.stdout + p.stderr
    events = [json.loads(l) for l in p.stdout.splitlines()]
    outputs = [Path(x) for x in events[-1]["outputs"]]
    assert outputs and all(x.exists() for x in outputs)
    for source, digest in before.items():
        assert hashlib.sha256(Path(source).read_bytes()).hexdigest() == digest
    return outputs


@pytest.mark.parametrize("source", ["png", "jpg", "webp", "bmp", "tiff", "gif"])
@pytest.mark.parametrize("target", ["png", "jpg", "webp", "bmp", "tiff", "gif", "pdf"])
def test_images(samples, tmp_path, source, target):
    from PIL import Image
    from pypdf import PdfReader

    output = execute(samples / f"imagem.{source}", target, tmp_path)[0]
    if target == "pdf":
        assert len(PdfReader(output).pages) == 1
    else:
        with Image.open(output) as im:
            assert im.size == (320, 240)
            im.load()


@pytest.mark.parametrize(
    "target", ["png", "jpg", "pdf-split", "pdf-extract", "pdf-compress"]
)
def test_pdf(samples, tmp_path, target):
    from pypdf import PdfReader

    output = execute(samples / "duas-paginas.pdf", target, tmp_path)[0]
    if output.is_dir():
        assert len(list(output.iterdir())) == 2
    else:
        assert len(PdfReader(output).pages) == 2


def test_merge_order_and_ranges(samples, tmp_path):
    from pypdf import PdfReader

    outputs = execute(
        [samples / "duas-paginas.pdf", samples / "duas-paginas.pdf"],
        "pdf-merge",
        tmp_path,
    )
    assert len(PdfReader(outputs[0]).pages) == 4
    request = dict(
        inputs=[str(outputs[0])],
        target="pdf-extract",
        outputDirectory=str(tmp_path),
        quality=85,
        pages="2-3",
        engines=ENGINES,
    )
    out = worker.convert(request)
    assert len(PdfReader(out[0]).pages) == 2


@pytest.mark.parametrize(
    ("source", "target"),
    [
        ("documento.docx", "pdf"),
        ("documento.docx", "odt"),
        ("documento.docx", "txt"),
        ("texto.txt", "docx"),
        ("planilha.xlsx", "pdf"),
        ("planilha.xlsx", "ods"),
        ("planilha.xlsx", "csv"),
        ("planilha.xlsx", "tsv"),
        ("tabela.csv", "xlsx"),
        ("slides.pptx", "pdf"),
        ("slides.pptx", "odp"),
    ],
)
def test_office(samples, tmp_path, source, target):
    from pypdf import PdfReader

    output = execute(samples / source, target, tmp_path)[0]
    if target == "pdf":
        assert len(PdfReader(output).pages) > 0
    elif target in ("txt", "csv", "tsv"):
        assert output.read_text(encoding="utf-8-sig").strip()
    else:
        assert output.read_bytes()[:2] == b"PK"


@pytest.mark.parametrize(
    "target", ["mp4", "mov", "mkv", "webm", "mp3", "wav", "flac", "aac", "ogg", "m4a"]
)
def test_media(samples, tmp_path, target):
    output = execute(samples / "video.mp4", target, tmp_path)[0]
    probe = Path(ENGINES["ffmpeg"]).with_name("ffprobe.exe")
    data = json.loads(
        worker.run([probe, "-v", "error", "-show_streams", "-of", "json", output])
    )
    assert any(s["codec_type"] == "audio" for s in data["streams"])
    if target in ("mp4", "mov", "mkv", "webm"):
        assert any(s["codec_type"] == "video" for s in data["streams"])


@pytest.mark.parametrize("target", ["woff", "woff2", "sfnt"])
def test_fonts(samples, tmp_path, target):
    from fontTools.ttLib import TTFont

    output = execute(samples / "fonte.ttf", target, tmp_path)[0]
    with TTFont(output) as font:
        assert font["maxp"].numGlyphs > 100


@pytest.mark.parametrize("target", ["pdf", "svg", "png"])
def test_cad(samples, tmp_path, target):
    output = execute(samples / "desenho.dxf", target, tmp_path)[0]
    if target == "svg":
        assert "<svg" in output.read_text()
    elif target == "pdf":
        assert output.read_bytes().startswith(b"%PDF")
    else:
        assert output.read_bytes().startswith(b"\x89PNG")


@pytest.mark.parametrize("target", ["zip", "7z"])
def test_archive_roundtrip(samples, tmp_path, target):
    output = execute([samples / "texto.txt", samples / "imagem.png"], target, tmp_path)[
        0
    ]
    folder = execute(output, "extract", tmp_path)[0]
    assert (folder / "texto.txt").read_bytes() == (samples / "texto.txt").read_bytes()
    assert (folder / "imagem.png").read_bytes() == (samples / "imagem.png").read_bytes()


def test_output_collision_and_accents(samples, tmp_path):
    folder = tmp_path / "ação"
    folder.mkdir()
    original = folder / "imagem.jpg"
    original.write_bytes(b"NAO SUBSTITUIR")
    result = execute(samples / "imagem.png", "jpg", folder)[0]
    assert result.name == "imagem (1).jpg"
    assert original.read_bytes() == b"NAO SUBSTITUIR"


def test_invalid_pdf_range(samples, tmp_path):
    with pytest.raises(ValueError):
        worker.pages_for("1-99", 2)
    with pytest.raises(ValueError):
        worker.pages_for("0", 2)
    with pytest.raises(ValueError):
        worker.pages_for("1;2", 2)


def test_archive_traversal_rejected():
    with pytest.raises(ValueError):
        worker.safe_archive_listing("----------\nPath = ../fora.txt")
    with pytest.raises(ValueError):
        worker.safe_archive_listing("----------\nPath = C:\\fora.txt")
    with pytest.raises(ValueError):
        worker.safe_archive_listing("----------\nPath = a\nSymbolic Link = ../fora")


def test_corrupt_file_fails_without_output(tmp_path):
    source = tmp_path / "danificado.png"
    source.write_bytes(b"not an image")
    out = tmp_path / "saida"
    result = subprocess.run(
        [sys.executable, str(ROOT / "worker" / "convert.py")],
        input=json.dumps(
            dict(
                inputs=[str(source)],
                target="webp",
                outputDirectory=str(out),
                engines=ENGINES,
            )
        ),
        text=True,
        capture_output=True,
        encoding="utf-8",
    )
    assert result.returncode == 1 and "error" in result.stdout
    assert not list(out.iterdir())
    assert source.read_bytes() == b"not an image"
